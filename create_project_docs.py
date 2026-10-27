from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path(r'C:\Users\srsad\eclipse-workspace\HighEducation')
NAVY='17324D'; TEAL='168C8C'; PALE='EAF2F6'; MID='5A6B78'; WHITE='FFFFFF'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell, text, bold=False, color=NAVY, size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
def setup(doc, title, subtitle):
    sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
    styles=doc.styles
    for name,size,color in [('Normal',10,'253746'),('Title',28,NAVY),('Heading 1',17,NAVY),('Heading 2',12,TEAL)]:
        s=styles[name]; s.font.name='Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
        if name.startswith('Heading'): s.font.bold=True
    styles['Normal'].paragraph_format.space_after=Pt(6)
    p=doc.add_paragraph(style='Title'); p.add_run(title)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(18); r=p.add_run(subtitle); r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(MID)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=footer.add_run('HIGHEDUCATION  |  PROJECT DOCUMENTATION'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(MID)
def para(doc,text): doc.add_paragraph(text)
def bullets(doc, items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],NAVY); set_cell_text(t.rows[0].cells[i],h,True,WHITE,9)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            set_cell_text(cells[i],str(v),False,NAVY,9); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(t.rows)%2==0: shade(cells[i],PALE)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Document 1: stakeholder overview
d=Document(); setup(d,'HighEducation Project Overview','Scholarship administration and CSR fund management')
para(d,'HighEducation is a web-based Java application that supports scholarship administration for KSS and corporate social responsibility (CSR) fund tracking for KSS and SSS. It brings approval, project, invoice, receipt, and vendor-payment records into one application, with institution-specific data kept in separate MySQL schemas.')
d.add_heading('Project at a glance',1)
table(d,['Item','Current project details'],[
('Application','HighEducation; Java web application packaged as a WAR'),('Institutions','KSS (scholarships and CSR) and SSS (CSR only)'),('Primary users','Global users can choose or switch institution; other users work in KSS'),('Technology','Java 17, Servlet 4, JSP, Maven, Tomcat 9, MySQL'),('Interface','Browser-based application with finance-focused CSR pages'),('Deployment','Railway build and deployment using Docker and Tomcat 9')],[1.5,5.7])
d.add_heading('What the application does',1)
d.add_heading('Scholarship administration',2)
para(d,'The KSS module provides scholarship portal functions for managing scholarship records and related documents. Scholarship pages are not available while a Global user is working in the SSS institution context.')
d.add_heading('CSR approvals and spending',2)
para(d,'The CSR workflow follows approved funding from management authorization through work orders, invoices, fund release letters (FRLs), receipts, and vendor payments. Users can review project positions and supporting financial details through dashboards, approval pages, ledgers, and printable statements.')
bullets(d,['Management approvals with statements, project/work-order positions, pending points, and print views.','Work orders, proforma and tax invoices, FRLs, and fund receipts.','Vendor advances, payment vouchers, deductions, and part payments against invoices.','Work-order debit and credit ledger with paid amounts and closing balance.','Vendor payment summaries showing bills, amounts paid, and balances due.','CSR dashboard with company and unit summaries, monthly receipts versus utilisation, and project status.','Company masters and company-aware approval, FRL, receipt, and print information.'])
d.add_heading('How institutions are separated',1)
para(d,'The application uses one login source and one MySQL server. KSS data is stored in the KSS schema; SSS CSR data is stored in SSS_CSR. A Global user selects an institution after login and can switch from the page header. Non-Global users remain in KSS. The application selects the corresponding schema for each request so records stay within the chosen institution.')
d.add_heading('Project status',1)
para(d,'The project status notes report the main CSR workflow, institution switching, company support, vendor payments, ledgers, and dashboards as implemented. Before deploying the current funding-company changes, the company schema migration must be applied in both KSS and SSS_CSR. The project notes also identify further work, including security improvements and payment-calculation rule updates.')
d.add_heading('Platform clarification',1)
para(d,'The inspected workspace is a Java web application, not a native Android application. It contains a Maven WAR project, Servlet controllers, JSP pages, and a Tomcat deployment configuration. These project details describe the application present in the workspace.')
d.add_heading('Purpose of this document',1)
para(d,'This overview is intended for project stakeholders who need a concise explanation of the application, its users, major functions, and current implementation status.')
d.save(OUT/'HighEducation_Project_Overview.docx')

# Document 2: technical specification
d=Document(); setup(d,'HighEducation Technical Specification','Architecture, modules, data, and operational notes')
para(d,'This specification summarizes the implementation described by the HighEducation source tree and project handover notes. It documents the current web application architecture and key operating rules for developers and maintainers.')
d.add_heading('1. System architecture',1)
table(d,['Layer','Implementation'],[
('Presentation','JSP pages, shared header/navigation, CSS including finance theme'),('Request handling','Java Servlets under com.Controller; ModuleFilter applies request context'),('Application configuration','AppConfig defines institution configuration and request-scoped selection'),('Data access','DAO classes and DBUtil provide MySQL access'),('Database','MySQL schemas KSS and SSS_CSR'),('Build/package','Maven WAR; Java source/target 17'),('Runtime/deployment','Tomcat 9; Docker-based Railway deployment')],[1.55,5.65])
d.add_heading('2. Login and institution context',1)
bullets(d,['Authentication is sourced from KSS.users.','After login, a Global user chooses KSS or SSS; the header supports switching between them.','Non-Global users are restricted to KSS.','ModuleFilter establishes the institution for each request; AppConfig exposes the selected configuration.','DBUtil connects to MySQL and selects the schema associated with the chosen institution.','KSS supports scholarship and CSR modules. SSS exposes CSR functions only; scholarship routes are redirected to CSR pages.'])
d.add_heading('3. Functional modules',1)
table(d,['Module','Responsibilities and notable behavior'],[
('Scholarships','KSS scholarship portal, scholarship records, and document-related operations.'),('Management approvals','Approval statements and project positions, tabs, pending points, company assignment, and A4 print view.'),('Work orders','Work-order management and per-order debit/credit ledger.'),('Invoices','Proforma and tax invoice handling; particulars field; proforma settlement and FRL requirement.'),('FRLs and receipts','Fund release letter claims and fund receipts linked to approvals/projects.'),('Vendor payments','Vendor advances, payment vouchers, deductions, part payments, and payment summaries.'),('CSR dashboard','Company/unit and project positions, monthly receipts versus utilisation, status/alerts, amount detail views.'),('Masters','Funding company records used by approvals and downstream CSR pages.')],[1.55,5.65])
d.add_heading('4. Financial workflow and rules',1)
para(d,'The implemented CSR sequence is Management Approval → Work Orders → Proforma or Tax Invoices → FRLs → Fund Receipts. Vendor payment tracking supports advances against work orders and one or more payment vouchers for a tax invoice. A work-order ledger treats bills as credits and payments/deductions as debits, then presents the closing balance.')
bullets(d,['Proforma invoices are not bills: they do not appear in the work-order ledger and do not receive payment vouchers.','A tax invoice can have multiple part-payment vouchers.','The current TDS sections are 194C and 194J.','Deletion and editing safeguards prevent changes that would invalidate linked downstream records.','Project notes identify payment-calculation behavior still requiring implementation: advance TDS/net handling, revised TDS and retention bases, retention release entries, and remarks for manual overrides. Confirm these rules with the project owner before changing calculations.'])
d.add_heading('5. Data and schema operations',1)
para(d,'The application uses identically named CSR tables in each institution schema. A CSR schema change generally needs to be applied to both KSS and SSS_CSR. New institutions are provisioned from the CSR database script described in HANDOVER.md.')
bullets(d,['KSS contains scholarship and CSR data; SSS_CSR contains CSR data only.','Funding-company support adds kss_csr_company and company_id on kss_csr_approval_letter.','The project status notes say Dump20261016\\KSS_csr_company.sql must be run once in each schema before deploying the company-aware code.','Payment-related database migrations and current schema state should be checked against PROJECT_STATUS.md and the SQL scripts before deployment.'])
d.add_heading('6. Key source areas',1)
table(d,['Area','Representative files'],[
('Institution and database context','AppConfig.java, ModuleFilter.java, InstitutionServlet.java, DBUtil.java, LoginServlet.java'),('CSR controllers','CsrDashboardServlet.java, CsrLetterServlet.java, CsrWorkOrderServlet.java, CsrInvoiceServlet.java, CsrFrlServlet.java, CsrFundServlet.java, CsrPaymentServlet.java, CsrCompanyServlet.java'),('Scholarship controllers and data','Scholarship servlets and DAO/Bean classes under com.Controller, com.DAO, and com.Bean'),('Views','CSR JSP pages, selectInstitution.jsp, shared header.jsp'),('Styling','css/finance.css, css/csrLetter.css'),('Build/deploy','pom.xml, Dockerfile')],[1.8,5.4])
d.add_heading('7. Build, run, and deployment',1)
bullets(d,['Local runtime described by the handover: Eclipse, Tomcat 9, Java 17, and Railway MySQL.','The Maven package is target/HighEducation-1.0-SNAPSHOT.war.','Railway builds the Docker image and deploys the application as ROOT.war on Tomcat 9.','Apply required SQL migrations before deploying code that depends on the new schema. Run CSR migrations in both institution schemas when applicable.','After source changes outside Eclipse, refresh the project, clean the build/server, restart Tomcat, and reload the browser to avoid stale files.'])
d.add_heading('8. Open maintenance items',1)
bullets(d,['Apply and verify the funding-company migration in both schemas before deploying related code.','Clean and test the application in Eclipse, then deploy.','Move the database password from source code to deployment configuration and rotate the exposed credential.','Hash user passwords and add CSRF protection to AJAX POST requests before wider use.','Review and implement the owner-defined vendor payment calculation rules.','Confirm the required KSS schema columns and FRL claim constraints.','Review placeholder menu items and other items listed in PROJECT_STATUS.md.'])
d.add_heading('9. Platform and scope',1)
para(d,'This repository does not contain a native Android project structure. The build configuration packages a Java web application as a WAR and targets Servlet/JSP on Tomcat. This specification therefore describes the existing HighEducation web system; a native Android client would be a separate project or future extension.')
d.save(OUT/'HighEducation_Technical_Specification.docx')
print('Created both DOCX files in', OUT)
